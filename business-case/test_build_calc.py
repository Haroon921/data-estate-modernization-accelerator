import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_PATH = Path(__file__).with_name("build_calc.py")
SPEC = importlib.util.spec_from_file_location("build_calc", SCRIPT_PATH)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError(f"Unable to load {SCRIPT_PATH}")
BUILD_CALC = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILD_CALC)


class BusinessCaseWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.output = Path(self.temp_dir.name) / "business_case.xlsx"
        BUILD_CALC.build_workbook(self.output)
        self.workbook = load_workbook(self.output, data_only=False)

    def tearDown(self):
        self.workbook.close()
        self.temp_dir.cleanup()

    def test_expected_sheets_are_present(self):
        self.assertEqual(
            self.workbook.sheetnames,
            [
                "Summary",
                "Inputs",
                "SQL_TCO",
                "Fabric_SKUs",
                "Synapse_to_Fabric",
                "Sensitivity",
            ],
        )

    def test_summary_contains_charts_and_validation_warning(self):
        summary = self.workbook["Summary"]
        self.assertEqual(len(summary._charts), 2)
        self.assertIn("COUNTIF", summary["A3"].value)
        self.assertEqual(summary["B6"].value, "=SQL_TCO!B12")

    def test_sensitivity_contains_three_scenarios_and_chart(self):
        sensitivity = self.workbook["Sensitivity"]
        self.assertEqual(
            [sensitivity[f"A{row}"].value for row in range(4, 7)],
            ["Downside", "Base", "Upside"],
        )
        self.assertEqual(len(sensitivity._charts), 1)
        self.assertTrue(sensitivity["H5"].value.startswith("="))

    def test_inputs_default_to_placeholder_status(self):
        inputs = self.workbook["Inputs"]
        statuses = [inputs[f"D{row}"].value for row in range(5, 17)]
        self.assertEqual(statuses, ["Placeholder"] * 12)


if __name__ == "__main__":
    unittest.main()
